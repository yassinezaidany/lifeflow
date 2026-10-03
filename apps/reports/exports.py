"""CSV / Excel exports of a user's entries and challenge statistics."""
from __future__ import annotations

import csv
import io

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.analytics.services.progress import ProgressEngine
from apps.challenges.models import Challenge
from apps.tracking.models import ChallengeEntry


def _entry_rows(user, challenge=None):
    qs = ChallengeEntry.objects.filter(user=user).select_related("challenge").prefetch_related("values__field").order_by("challenge__name", "date")
    if challenge is not None:
        qs = qs.filter(challenge=challenge)
    keys: list[str] = []
    rows = []
    for e in qs:
        values = {v.field.label: v.value for v in e.values.all()}
        for k in values:
            if k not in keys:
                keys.append(k)
        rows.append((e, values))
    header = ["Challenge", "Date", *keys, "Note", "Source"]
    data = [[e.challenge.name, e.date.isoformat(), *[_cell(values.get(k)) for k in keys], e.note, e.source] for e, values in rows]
    return header, data


def _stats_rows(user):
    challenges = list(ProgressEngine.prefetch(Challenge.objects.filter(user=user)))
    results = ProgressEngine(user).evaluate_many(challenges)
    header = ["Challenge", "Status", "Goal", "Start", "End", "Target to date", "Actual", "Expected", "Progress %",
              "Gap", "Completion %", "Current streak", "Best streak", "Unit"]
    data = []
    for c in challenges:
        r = results[c.pk]
        data.append([c.name, c.status, r.goal_description, c.start_date.isoformat(), c.end_date.isoformat() if c.end_date else "",
                     round(r.goal, 2), round(r.actual, 2), _round(r.expected), _round(r.progress, 1), _round(r.gap),
                     _round(r.completion_rate, 1), r.current_streak, r.best_streak, r.unit or ("min" if r.is_duration else "")])
    return header, data


def _cell(value):
    if value is True:
        return "yes"
    if value is False:
        return "no"
    return "" if value is None else value


def _round(v, n=2):
    return None if v is None else round(v, n)


def _safe(value):
    """Neutralise spreadsheet formula injection."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def export_csv(user, kind: str, challenge=None) -> HttpResponse:
    header, data = _stats_rows(user) if kind == "statistics" else _entry_rows(user, challenge)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    for row in data:
        writer.writerow([_safe(v) for v in row])
    response = HttpResponse("﻿" + buf.getvalue(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="lifeflow-{kind}.csv"'
    return response


def export_xlsx(user) -> HttpResponse:
    wb = Workbook()
    for index, (title, (header, data)) in enumerate((("Statistics", _stats_rows(user)), ("Entries", _entry_rows(user)))):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4F46E5")
        for row in data:
            ws.append([_safe(v) for v in row])
        for column in ws.columns:
            width = max(len(str(c.value or "")) for c in column)
            ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 40)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="lifeflow-export.xlsx"'
    return response
